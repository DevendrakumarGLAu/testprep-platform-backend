"""Source parsers. Each returns (headers, rows, warnings); rows are dicts of header -> text."""
import csv
import io
import re
from decimal import Decimal

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024


class ParseError(Exception):
    """A problem with the uploaded file that the user can fix."""


def _unique_headers(headers):
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(headers):
        name = (h or "").strip() or f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, Decimal) and v == v.to_integral():
        return str(int(v))
    return str(v).strip()


def _tabulate(table: list[list]):
    """First non-empty row is the header; fully blank rows are dropped."""
    rows = [r for r in table if any(_cell(c) for c in r)]
    if not rows:
        raise ParseError("The file is empty.")
    headers = _unique_headers([_cell(c) for c in rows[0]])
    data = []
    for r in rows[1:]:
        data.append({h: _cell(r[i]) if i < len(r) else "" for i, h in enumerate(headers)})
    if not data:
        raise ParseError("The file has a header row but no questions.")
    if len(data) > MAX_ROWS:
        raise ParseError(f"Too many rows ({len(data)}). Split the file into batches of at most {MAX_ROWS}.")
    return headers, data


def parse_csv(raw: bytes):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ParseError("The CSV must be saved as UTF-8. In Excel choose 'CSV UTF-8 (Comma delimited)'.")
    first = text.splitlines()[0] if text.strip() else ""
    delimiter = max([",", ";", "\t"], key=first.count) if first else ","
    try:
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as e:
        raise ParseError(f"The CSV could not be read: {e}")
    headers, rows = _tabulate(table)
    return headers, rows, []


def parse_xlsx(raw: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)  # data_only: never evaluates formulas
    except Exception:
        raise ParseError("This is not a valid .xlsx file.")
    ws = wb.worksheets[0]
    warnings = []
    if len(wb.worksheets) > 1:
        warnings.append(f"Only the first sheet ('{ws.title}') was read.")
    table = []
    for row in ws.iter_rows(values_only=True):
        table.append(list(row))
        if len(table) > MAX_ROWS + 50:
            break
    wb.close()
    headers, rows = _tabulate(table)
    return headers, rows, warnings


def parse_xls(raw: bytes):
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception:
        raise ParseError("This is not a valid .xls file.")
    sheet = book.sheet_by_index(0)
    warnings = [f"Only the first sheet ('{sheet.name}') was read."] if book.nsheets > 1 else []
    table = [sheet.row_values(i) for i in range(sheet.nrows)]
    headers, rows = _tabulate(table)
    return headers, rows, warnings


# --- PDF -------------------------------------------------------------------------------------

QUESTION_START = re.compile(r"^\s*(?:Q(?:uestion)?\.?\s*)?(\d{1,4})\s*[.)\]:]\s+(.*\S)", re.I)
OPTION_LINE = re.compile(r"^\s*\(?([A-Ja-j])\s*[.)\]:]\s+(.*\S)")
ANSWER_LINE = re.compile(r"^\s*(?:correct\s+answer|correct\s+option|answer|ans|key)\s*[:\-]\s*(.*\S)", re.I)
EXPLANATION_LINE = re.compile(r"^\s*(?:explanation|solution|rationale)\s*[:\-]\s*(.*)", re.I)
INLINE_OPTIONS = re.compile(r"(?:^|\s)\(?([A-D])[.)]\s")


def parse_pdf_text(text: str):
    """
    Turns extracted PDF text into the same table the CSV/Excel parsers produce.
    Recognises 'Q1.' / '1)' question starts, 'A.' / '(b)' option lines, 'Answer:' and 'Explanation:'.
    """
    questions = []
    cur = None
    field = None  # which part continuation lines belong to
    for line in text.splitlines():
        if not line.strip():
            continue
        m_ans = ANSWER_LINE.match(line)
        m_exp = EXPLANATION_LINE.match(line)
        m_opt = OPTION_LINE.match(line)
        m_q = QUESTION_START.match(line)

        if cur and m_ans:
            cur["answer"] = m_ans.group(1).strip()
            field = "answer"
        elif cur and m_exp:
            cur["explanation"] = m_exp.group(1).strip()
            field = "explanation"
        elif cur and m_opt and (field in (None, "question") or (field or "").startswith("option:")):
            letter = m_opt.group(1).upper()
            if letter not in cur["options"]:
                cur["options"][letter] = m_opt.group(2).strip()
                field = f"option:{letter}"
            else:
                _append(cur, field, line)
        elif m_q and (cur is None or cur["options"] or field in ("answer", "explanation") or not cur["text"]):
            cur = {"number": m_q.group(1), "text": m_q.group(2).strip(), "options": {}, "answer": "", "explanation": ""}
            questions.append(cur)
            field = "question"
        elif cur:
            _append(cur, field, line)

    if not questions:
        raise ParseError(
            "No questions were recognised. Questions should look like '1. Question text', options like "
            "'A. option', followed by 'Answer: B'."
        )
    headers = ["question", "option_a", "option_b", "option_c", "option_d", "option_e", "option_f",
               "correct_answer", "explanation"]
    rows = []
    warnings = []
    for q in questions[:MAX_ROWS]:
        row = {"question": q["text"], "correct_answer": q["answer"], "explanation": q["explanation"]}
        for letter in "ABCDEF":
            row[f"option_{letter.lower()}"] = q["options"].get(letter, "")
        rows.append(row)
    if len(questions) > MAX_ROWS:
        warnings.append(f"Only the first {MAX_ROWS} questions were read.")
    if not any(r["correct_answer"] for r in rows):
        warnings.append("No 'Answer:' lines were found. Set the correct answer for each question during review.")
    return headers, rows, warnings


def _append(cur, field, line):
    line = line.strip()
    if field == "question" or field is None:
        cur["text"] += " " + line
    elif field == "answer":
        cur["answer"] += " " + line
    elif field == "explanation":
        cur["explanation"] += " " + line
    elif field and field.startswith("option:"):
        letter = field.split(":")[1]
        cur["options"][letter] += " " + line


def parse_pdf(raw: bytes):
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ParseError("This PDF is password protected.")
        pages = [(p.extract_text() or "") for p in reader.pages[:200]]
    except ParseError:
        raise
    except Exception:
        raise ParseError("This is not a readable PDF file.")
    text = "\n".join(pages)
    if len(text.strip()) < 10:
        raise ParseError(
            "No text could be extracted. Scanned or image-only PDFs are not supported (OCR is not available)."
        )
    headers, rows, warnings = parse_pdf_text(text)
    warnings.insert(0, "PDF layouts vary. Check every question carefully before importing.")
    return headers, rows, warnings


def detect_source(filename: str) -> str:
    name = filename.lower()
    if name.endswith(".csv"):
        return "csv"
    if name.endswith((".xlsx", ".xls")):
        return "excel"
    if name.endswith(".pdf"):
        return "pdf"
    raise ParseError("Unsupported file type. Upload a .csv, .xlsx, .xls or .pdf file.")


def parse_file(filename: str, raw: bytes):
    if len(raw) > MAX_BYTES:
        raise ParseError("The file is larger than 5 MB.")
    if not raw:
        raise ParseError("The file is empty.")
    source = detect_source(filename)
    if source == "csv":
        return source, *parse_csv(raw)
    if source == "pdf":
        return source, *parse_pdf(raw)
    if filename.lower().endswith(".xls"):
        return source, *parse_xls(raw)
    return source, *parse_xlsx(raw)
